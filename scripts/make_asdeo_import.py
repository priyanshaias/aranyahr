"""Makes private/AS_import.csv and private/DEO_import.csv: the Aranya Sathis (PDL List.xlsx, sheet "Final Remuneration (1555)")
and the DEOs (DEO.xlsx) in the columns of the "Aranya Sathi" and "DEO" tabs of the AranyaHR · AS+DEO sheet.
Append each to its tab (File > Import > Upload > "Append to current sheet"; see the README). Real names: the files stay in private/.
Run: python3 scripts/make_asdeo_import.py raw_data   (needs openpyxl)"""
import sys, os, re, csv, json, secrets, datetime
from collections import Counter
from openpyxl import load_workbook

SRC = sys.argv[1]
OUT = os.path.join(os.path.dirname(__file__), '..', 'private')
# the columns of the Aranya Sathi and DEO tabs (Cdlbs.gs COLS, without "As in the return")
COLS = ['ID', 'Circle', 'Division', 'Name', 'Mode of engagement', 'Agency name', 'Engagement date', 'District', 'Range (or Division HQ)', 'Beat',
        'Place of work', 'Name of the place', 'Main duty', 'Other duties', 'Duty details', 'Monthly pay (Rs)', 'Remarks', 'Verification', 'Status',
        'Problem', 'Removal request', 'Last updated on', 'Last updated by']
DIRECTORATE = ('APM', 'APM Circle HQ')   # the Directorate's own office: project DEOs and "APM / Direction"

# PDL list: circle -> AranyaHR circle; (circle, division) -> AranyaHR division ("Circle HQ" -> "<circle> Circle HQ")
PDL_CIRCLE = {'P&G': 'Parks & Gardens', 'SC (N)': 'Soil Con (N)', 'Wildlife (N)': 'Wild Life North'}
PDL_DIV = {'Direction': None, 'BTR (East)': 'B.T.R (East)', 'BTR (West)': 'B.T.R (West)', 'SFTI, Hijli': 'SFTI Hizli', 'WBFS, Dow Hill': 'WB Forest School',
  'Kurseong': 'Kurseong Dn.', 'Monitoring (N)': 'Monitoring North', 'Monitoring (S)': 'Monitoring South', 'Jalpaiguri SF': 'Jalpaiguri S.F.', 'Siliguri SF': 'Siliguri S.F',
  'Cooch Behar': 'Coochbehar', 'P&G (N)': 'Parks & Gardens North', '24 Parganas North': '24 Parganas (N)', '24 Parganas South': '24 Parganas (S)',
  'Nadia Murshidabad': 'Nadia-Murshidabad', 'Jalpaiguir SC': 'Jalpaiguri SC', 'STR': 'S.T.R', 'Extension Forestry': 'Extn. Forestry', 'Kangsabati North': 'Kangsabati (N)',
  'WL HQ': 'WL HQ Division', 'WP (N)': 'WP North', 'WP (South)-I': 'WP South I', 'WP (South)-II': 'WP South II'}
# DEO list: office -> (circle, division)
DEO_DIV = {
  'West Bengal Forest & Bio Diversity Conservation Project under Directorate of Forests': DIRECTORATE,
  'West Bengal State Forests Devlopment Agency under Directorate of Forests': DIRECTORATE,
  'West Bengal Compansatory Afforestation Management and Planning Authority under Directorate of Forests': DIRECTORATE,
  'Integrated Watershed Management Project': DIRECTORATE, 'O/o the PCCF, HoFF': DIRECTORATE,
  'O/o the DFO,Bankura (S)': ('Central', 'Bankura (S)'), 'O/o the DFO, Kangsabati North': ('South West', 'Kangsabati (N)'), 'O/o the DFO,Burdwan': ('South East', 'Burdwan'),
  'O/o the CF WP &GIS': ('WP & GIS', 'WP & GIS Circle HQ'), 'O/o the CF,Park & Gardens': ('Parks & Gardens', 'Parks & Gardens Circle HQ'),
  'O/o the DFD BTR (East) ,Alipurduar': ('BTR', 'B.T.R (East)'), 'O/o the DFD BTR (West) ,Alipurduar': ('BTR', 'B.T.R (West)'), 'O/o the F.D & CCF, BTR, Alipurduar': ('BTR', 'BTR Circle HQ'),
  'O/o the DFD, STR': ('STR', 'S.T.R'), 'O/o the DFO,Kurseong': ('Hill', 'Kurseong Dn.'), 'CCF,SOUTH-EAST,DURGAPUR': ('South East', 'South East Circle HQ'),
  'O/o the DFO, Birbhum': ('South East', 'Birbhum'), 'O/o the DFO,Durgapur': ('South East', 'Durgapur'), 'O/o the DFO, Howrah': ('Parks & Gardens', 'Howrah'),
  'O/o the DFO,Publicity Division': ('APM', 'Publicity'), 'O/o the CF,Devlopment': ('Development', 'Development Circle HQ'), 'O/o the PCCF,WL': ('WL HQ', 'WL HQ Division'),
  'O/o the CCF, WL (N)': ('Wild Life North', 'Wild Life North Circle HQ'), 'O/o the DFO, Gorumara WL': ('Wild Life North', 'Gorumara WL'),
  'O/o the DFO, Monitoring (S)': ('Monitoring', 'Monitoring South'), 'O/o the DFO, Purulia': ('South West', 'Purulia'), 'O/o the DFO Purulia': ('South West', 'Purulia'),
  'O/o the DFO, Siliguri SF': ('North West', 'Siliguri S.F'), 'O/o the DFO,Jalguri': ('Northern', 'Jalpaiguri'), 'O/o the DFO,Medinipur': ('Western', 'Medinipur'),
  'O/o the DFO,Purba Medinipur': ('Western', 'Purba Medinipur'),
}

clean = lambda s: re.sub(r'\s+', ' ', str(s if s is not None else '')).strip()
def day(v):
    """A date cell or text (03.06.2019, 3/6/2019) -> dd-mm-yyyy (the list is day first)."""
    if isinstance(v, datetime.datetime): return v.strftime('%d-%m-%Y')
    m = re.match(r'^(\d{1,2})[./-](\d{1,2})[./-](\d{4})$', clean(v))
    return '%02d-%02d-%s' % (int(m[1]), int(m[2]), m[3]) if m else ''
ids = set()
def new_id(pre):
    while True:
        x = pre + '-' + secrets.token_hex(4).upper()
        if x not in ids: ids.add(x); return x
NOW = datetime.datetime.now().strftime('%d-%m-%Y %H:%M:%S')

def write(name, people, cadre, source):
    """One CSV with no heading row (it goes under the tab's own headings), and a possible-double-entry flag like the BS import."""
    key = lambda o: (o['Division'], o['Name'].upper(), o.get('Engagement date', ''))
    twice = Counter(key(o) for o in people)
    with open(os.path.join(OUT, name), 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        for o in people:
            o = dict(o, ID=new_id(cadre), **{'Last updated by': 'Import: ' + source})
            if twice[key(o)] > 1: o['Problem'] = 'Possible double entry: same name and engagement date appear twice in this division.'
            w.writerow([o.get(h, '') for h in COLS])
    by = Counter((o['Circle'], o['Division']) for o in people)
    print('%s: %d people in %d divisions -> private/%s · possible double entries: %d' % (cadre, len(people), len(by), name, sum(1 for o in people if twice[key(o)] > 1)))
    return by

os.makedirs(OUT, exist_ok=True)
divs = set()   # AranyaHR's divisions, to check every row lands in one (private/lists.json: the API's "lists" reply)
if os.path.exists(os.path.join(OUT, 'lists.json')):
    for c in json.load(open(os.path.join(OUT, 'lists.json')))['circles']:
        for d in c['divisions']: divs.add((c['circle'], d))

# ---------------------------------------------------------------- Aranya Sathi (PDL)
ws = load_workbook(os.path.join(SRC, 'PDL List.xlsx'), read_only=True, data_only=True)['Final Remuneration (1555)']
rows = list(ws.iter_rows(values_only=True))
hi = next(i for i, r in enumerate(rows) if r and r[3] and 'Name of the listed' in str(r[3]))
head = [clean(h).lower() for h in rows[hi]]
col = lambda words: next(i for i, h in enumerate(head) if all(w in h for w in words))
C = {'circle': col(['circle']), 'div': col(['division']), 'name': col(['name of the listed']), 'cat': col(['category']), 'go': col(['engaged vide']),
     'gosl': col(['sl. no. in list']), 'dob': col(['date of birth']), 'dor': col(['date of retirement']), 'st': col(['retired']), 'doe': col(['date of engagement']),
     'pay': col(['remuneration as on 01/01/2026'])}
AS, odd = [], Counter()
for r in rows[hi + 2:]:
    name = clean(r[C['name']]) if r and len(r) > C['name'] else ''
    if not name or name.lower() == 'nil': continue
    c0, d0 = clean(r[C['circle']]), clean(r[C['div']])
    circle = PDL_CIRCLE.get(c0, c0)
    div = PDL_DIV.get(d0, d0)
    if d0 == 'Direction': circle, div = DIRECTORATE
    elif d0 == 'Circle HQ': div = circle + ' Circle HQ'
    st = clean(r[C['st']])
    notes = ['Category ' + clean(r[C['cat']])] if clean(r[C['cat']]) else []
    if clean(r[C['go']]): notes.append('G.O. ' + clean(r[C['go']]) + (' (sl. ' + clean(r[C['gosl']]) + ')' if clean(r[C['gosl']]) else ''))
    if day(r[C['dob']]): notes.append('born ' + day(r[C['dob']]))
    if day(r[C['dor']]): notes.append('retires ' + day(r[C['dor']]))
    o = {'Circle': circle, 'Division': div, 'Name': name, 'Mode of engagement': 'Direct', 'Engagement date': day(r[C['doe']]), 'Status': 'Working'}
    pay = r[C['pay']]
    if isinstance(pay, (int, float)): o['Monthly pay (Rs)'] = str(int(round(pay)))
    if st == 'Died':   # as when a division marks someone deceased: waits for the admin to confirm the removal
        o['Status'] = 'Removal pending'
        o['Removal request'] = json.dumps({'reason': 'Deceased', 'remark': 'Marked "Died" in the PDL list', 'by': 'Import: PDL list', 'on': NOW})
        notes.append('deceased (PDL list)')
    elif st != 'In Service':
        odd[st or '(blank)'] += 1
        notes.append('status in the PDL list: ' + (st or 'blank') + ' (please check)')
    o['Remarks'] = '; '.join(notes)
    AS.append(o)
by_as = write('AS_import.csv', AS, 'AS', 'PDL List, Final Remuneration (1555)')
print('  status: %d working, %d deceased (removal pending), unclear: %s' % (sum(1 for o in AS if o['Status'] == 'Working'), sum(1 for o in AS if o['Status'] != 'Working'), dict(odd)))

# ---------------------------------------------------------------- DEO
ws = load_workbook(os.path.join(SRC, 'DEO.xlsx'), read_only=True, data_only=True).worksheets[0]
DEO, unknown = [], Counter()
for r in ws.iter_rows(values_only=True):
    if not r or not isinstance(r[0], (int, float)) or not clean(r[1]): continue
    office = clean(r[4])
    if office not in DEO_DIV: unknown[office] += 1; continue
    circle, div = DEO_DIV[office]
    desig, project = clean(r[2]), clean(r[9])
    notes = [desig, office]
    if project and project not in office: notes.append('project ' + project)
    if isinstance(r[19], (int, float)): notes.append('pay Rs %d on 31-10-2020' % r[19])
    software = 'software' in desig.lower()
    DEO.append({'Circle': circle, 'Division': div, 'Name': clean(r[1]), 'Engagement date': day(r[18]), 'Main duty': 'Others',
                'Duty details': 'Software support and e-governance work' if software else 'Data entry and e-governance work',
                'Remarks': '; '.join(notes), 'Status': 'Working'})
by_deo = write('DEO_import.csv', DEO, 'DEO', 'DEO.xlsx, Final list as on 22.04.2026')
if unknown: print('  DEO offices with no division (not written):', dict(unknown))

if divs:
    bad = [d for d in sorted(set(by_as) | set(by_deo)) if d not in divs]
    print('divisions not in AranyaHR:', bad or 'none')
print('AS by division:', json.dumps({d: n for (c, d), n in sorted(by_as.items())}, ensure_ascii=False))
print('DEO by division:', json.dumps({d: n for (c, d), n in sorted(by_deo.items())}, ensure_ascii=False))
